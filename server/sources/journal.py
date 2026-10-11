"""The manual trade journal workbook (`1.1 Trade Journal2.xlsx`): reading it, and mirroring it into the trades table.

A trade is a row with a date in column B and a pair in column E. Amounts stay in the journal's own currency (USDT).
"""
from datetime import datetime

DEFAULT_SHEET = "Trading Journal Fianl Test"     # spelled as in the workbook


def cell(row, letter):
    """Value by column letter(s); the journal has 29 columns (A..AC)."""
    n = 0
    for ch in letter:
        n = n * 26 + ord(ch) - 64
    return row[n - 1] if n <= len(row) else None


def clean(v):
    return str(v).strip() if v not in (None, "") else None


def link(v):
    """A chart link, or None for anything that is not one (the sheet's header rows say 'Pic')."""
    v = clean(v)
    return v if v and v.lower().startswith(("https://", "http://")) else None


def number(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def parse(path, sheet=DEFAULT_SHEET):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    trades = []
    for n, r in enumerate(wb[sheet].iter_rows(min_row=1, max_col=29, values_only=True), start=1):
        when, pair = cell(r, "B"), clean(cell(r, "E"))
        if not isinstance(when, datetime) or not pair:
            continue
        outcome = "win" if cell(r, "V") else "loss" if cell(r, "W") else "be" if cell(r, "X") else None
        # Profit (Y) and Loss (Z) are kept whatever the outcome: break-even trades carry small amounts,
        # and the sheet's own PNL is the sum of Y minus the sum of Z.
        profit, lost = number(cell(r, "Y")), number(cell(r, "Z"))
        pnl = None if profit is None and lost is None and outcome is None else (profit or 0) - abs(lost or 0)
        trades.append({"source_row": n, "traded_on": when.strftime("%Y-%m-%d"), "session": clean(cell(r, "D")),
                       "pair": pair, "side": clean(cell(r, "F")), "setup": clean(cell(r, "G")),
                       "exit_plan": clean(cell(r, "P")), "exit_actual": clean(cell(r, "Q")),
                       "grade": clean(cell(r, "U")), "outcome": outcome, "pnl": pnl,
                       "balance": number(cell(r, "AB")), "risk": number(cell(r, "AC")), "note": clean(cell(r, "T")),
                       # War Zone (I..K), Entry (L..O), Exit's "100 of A?" (R) and the three chart links (H, O, S)
                       "pic_setup": link(cell(r, "H")), "poi": clean(cell(r, "I")), "signal": clean(cell(r, "J")),
                       "divergence": clean(cell(r, "K")), "confirm": clean(cell(r, "L")), "fib_time": clean(cell(r, "M")),
                       "fib": clean(cell(r, "N")), "pic_entry": link(cell(r, "O")), "exit_a": clean(cell(r, "R")),
                       "pic_exit": link(cell(r, "S")), "rr": number(cell(r, "AA"))})
    wb.close()
    return trades


def mirror(conn, trades):
    """Make the trades table exactly the sheet's trades: whatever is not in the sheet any more is removed. One transaction.
    Returns (added, removed) counts, comparing trades by what they say (date, pair, side, setup, P&L)."""
    def fingerprint(t):
        return (t["traded_on"], t["pair"], t["side"], t["setup"], round(t["pnl"] or 0, 2))
    old = [fingerprint(dict(r)) for r in conn.execute("SELECT traded_on, pair, side, setup, pnl FROM trades")]
    new = [fingerprint(t) for t in trades]
    gone = list(old)
    added = 0
    for f in new:
        if f in gone:
            gone.remove(f)
        else:
            added += 1
    cols = list(trades[0])
    conn.execute("DELETE FROM trades")
    conn.executemany(f"INSERT INTO trades({', '.join(cols)}) VALUES ({', '.join('?' for _ in cols)})", [[t[c] for c in cols] for t in trades])
    conn.commit()
    return added, len(gone)
