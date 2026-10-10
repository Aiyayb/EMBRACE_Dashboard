"""Family Statement, read from the ledger file Maki maintains (ADR 0001).

It is the extended family's joint account, not Best's money, so it is always
Display-only. Monthly totals are computed from the transactions here because
the workbook's own summary sheets go stale.
"""
import os
from datetime import date, datetime

_CACHE = {"key": None, "value": None}


def _iso(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v)[:10] if v else None


def ledger(xlsx_path):
    """{"available", "reason", "balance", "transactions": [...], "months": [...]}"""
    if not os.path.exists(xlsx_path):
        return {"available": False, "reason": f"ledger not found at {xlsx_path}", "transactions": [], "months": []}
    key = (xlsx_path, os.path.getmtime(xlsx_path))
    if _CACHE["key"] == key:
        return _CACHE["value"]
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"available": False, "reason": "openpyxl is not installed (pip install openpyxl)",
                "transactions": [], "months": []}
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    try:
        if "Statement" not in wb.sheetnames:
            return {"available": False, "reason": "sheet 'Statement' not found", "transactions": [], "months": []}
        rows = list(wb["Statement"].iter_rows(min_row=2, max_col=11, values_only=True))
    finally:
        wb.close()
    transactions, months, balance = [], {}, 0.0
    for r in rows:
        if r[0] is None or r[1] is None:
            continue
        income, expense = float(r[5] or 0), float(r[6] or 0)
        balance += income - expense
        on = _iso(r[1])
        transactions.append({"id": r[0], "date": on, "type": r[2], "category": r[3], "detail": r[4],
                             "income": income, "expense": expense, "balance": round(balance, 2),
                             "recorded_by": r[8], "note": r[9]})
        m = months.setdefault(on[:7], {"month": on[:7], "income": 0.0, "expense": 0.0, "count": 0})
        # the opening balance is carried in, not earned that month
        if r[3] != "ยอดยกมา":
            m["income"] += income
        m["expense"] += expense
        m["count"] += 1
    for m in months.values():
        m["net"] = round(m["income"] - m["expense"], 2)
    value = {"available": True, "reason": None, "balance": round(balance, 2),
             "last_date": transactions[-1]["date"] if transactions else None,
             "transactions": transactions, "months": sorted(months.values(), key=lambda m: m["month"])}
    _CACHE.update(key=key, value=value)
    return value
