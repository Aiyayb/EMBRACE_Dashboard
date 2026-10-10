import os
import sys
import tempfile
import unittest
from datetime import datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "server"))
sys.path.insert(0, os.path.join(ROOT, "scripts"))

import db  # noqa: E402
import metrics  # noqa: E402

try:
    from openpyxl import Workbook
except ImportError:                      # the importers need openpyxl; skip where it is absent
    Workbook = None


def build_workbook(path, year_start_usdt=None):
    """A tiny workbook with the same block layout as `4. Fund Management.xlsx`."""
    wb = Workbook()
    stock = wb.active
    stock.title = "Stock Portfolio"
    stock["G4"], stock["H4"], stock["J4"] = "Year 2026", 1200, 700000      # baseline row: skipped
    for row, (when, set_index, value) in enumerate([(datetime(2025, 12, 25), 1200, 700000),
                                                    (datetime(2026, 1, 26), 1260, 730000),
                                                    (datetime(2026, 2, 26), 1300, 760000),
                                                    (datetime(2026, 3, 26), 1280, 750000)], start=5):
        stock[f"G{row}"], stock[f"H{row}"], stock[f"J{row}"] = when, set_index, value
    stock["G9"] = datetime(2026, 4, 26)                                   # future row, no value: skipped

    trading = wb.create_sheet("Trading Portfolio")
    if year_start_usdt is not None:                                       # the row that opens the year
        trading["G4"], trading["H4"], trading["J4"], trading["K4"] = "Year 2026", 79000, year_start_usdt, year_start_usdt * 33
    for row, (when, btc, usdt) in enumerate([(datetime(2026, 1, 26), 80000, 1000),
                                             (datetime(2026, 2, 26), 82000, 1100),
                                             (datetime(2026, 3, 26), 81000, 1200)], start=5):
        trading[f"G{row}"], trading[f"H{row}"], trading[f"J{row}"], trading[f"K{row}"] = when, btc, usdt, usdt * 33

    alloc = wb.create_sheet("Asset Allocation(Q)")
    # B label, C total, E EMC, J stock, L Saving, N Trading, P Bot
    for row, (label, total, emc, stk, saving, trade, bot) in enumerate(
            [("Q4 - 2025", 780000, 50000, 700000, 0, 30000, 0), ("Q1 - 2026", 839600, 50000, 750000, 0, 39600, 0),
             ("Q2 - 2026", 0, None, None, None, None, None)], start=39):
        for col, v in zip("BCEJLNP", (label, total, emc, stk, saving, trade, bot)):
            alloc[f"{col}{row}"] = v

    plan = wb.create_sheet("Fin Plan(Y)")
    plan["C29"] = "TARGET"
    for col, v in zip("GHJLNP", (2026, 100000, 900000, 100000, 60000, 50000)):
        plan[f"{col}31"] = v
    for col, month in zip("HIJKLMNOPQRS", range(1, 13)):
        plan[f"{col}53"] = datetime(2026, month, 26)
    plan["F54"], plan["G54"], plan["H54"], plan["I54"], plan["J54"], plan["K54"] = "Stock", 14000, 15000, 7000, -5000, 9999
    plan["F55"], plan["H55"] = "Trading", 3000
    plan["F57"], plan["H57"] = "EMC", 2000
    wb.save(path)


@unittest.skipIf(Workbook is None, "openpyxl not installed")
class ImportTest(unittest.TestCase):
    def setUp(self):
        import import_fund_management as imp
        import reconcile
        self.imp, self.reconcile = imp, reconcile
        self.tmp = tempfile.TemporaryDirectory()
        self.xlsx = os.path.join(self.tmp.name, "fund.xlsx")
        build_workbook(self.xlsx)
        self.conn = db.connect(os.path.join(self.tmp.name, "t.db"))
        db.init(self.conn)

    def tearDown(self):
        self.conn.close()
        self.tmp.cleanup()

    def test_parse_finds_rows_by_content(self):
        p = self.imp.parse(self.xlsx)
        self.assertEqual([s["as_of"] for s in p["stock"]], ["2025-12-25", "2026-01-26", "2026-02-26", "2026-03-26"])
        self.assertEqual(len(p["trading"]), 3)
        self.assertEqual([q["label"] for q in p["quarters"]], ["Q4 - 2025", "Q1 - 2026"])   # empty quarter skipped
        # the April contribution is beyond the last stock Snapshot, so it is not imported
        self.assertEqual(sorted((f["account"], f["on_date"], f["amount"]) for f in p["flows"]),
                         [("OKX Trading", "2026-01-26", 3000.0), ("Stocks", "2026-01-26", 15000.0),
                          ("Stocks", "2026-02-26", 7000.0), ("Stocks", "2026-03-26", -5000.0)])
        self.assertTrue(any("EMC" in w for w in p["warnings"]))
        self.assertEqual({g["name"]: g["goal"] for g in p["goals"]},
                         {"OnTheWall": 100000, "Stock": 900000, "Trading": 100000, "Saving": 60000, "EMC": 50000})

    def test_import_reconciles_and_states_return_net_of_flows(self):
        p = self.imp.parse(self.xlsx)
        self.imp.apply(self.conn, p)
        self.assertEqual(db.get_setting(self.conn, "flows_recorded_from"), "2026-01")
        self.assertTrue(all(r["ok"] for r in self.reconcile.compare(self.conn, p)))
        jan = next(m for m in metrics.monthly(self.conn) if m["month"] == "2026-01")
        # stocks +30,000 with 15,000 deposited; trading 33,000 from 30,000 with 3,000 deposited
        self.assertEqual(jan["return_thb"], 15000)
        self.assertEqual(jan["trading_pnl"], 0)

    def test_year_opening_trading_balance_overrides_the_quarterly_figure(self):
        build_workbook(self.xlsx, year_start_usdt=250)
        p = self.imp.parse(self.xlsx)
        self.imp.apply(self.conn, p)
        dec = next(m for m in metrics.monthly(self.conn) if m["month"] == "2025-12")
        self.assertEqual(dec["trading"], 8250)                # not the 30,000 in Asset Allocation(Q)
        jan = next(m for m in metrics.monthly(self.conn) if m["month"] == "2026-01")
        self.assertEqual(jan["trading_pnl"], 33000 - 8250 - 3000)
        # the workbook's own Q4 total still holds the old figure, and reconcile says so
        q4 = next(r for r in self.reconcile.compare(self.conn, p) if r["label"] == "Q4 - 2025")
        self.assertFalse(q4["ok"])
        self.assertEqual(q4["diff"], 8250 - 30000)
        self.assertEqual(self.conn.execute(
            "SELECT value FROM benchmarks WHERE name='BTC' AND as_of='2025-12-31'").fetchone()["value"], 79000)

    def test_reimport_keeps_figures_entered_in_the_dashboard(self):
        p = self.imp.parse(self.xlsx)
        self.imp.apply(self.conn, p)
        stocks = self.conn.execute("SELECT id FROM accounts WHERE name='Stocks'").fetchone()["id"]
        self.conn.execute("UPDATE snapshots SET value_thb=1, value_native=1, entered_by='best' "
                          "WHERE account_id=? AND as_of='2026-03-26'", (stocks,))
        self.conn.commit()
        self.imp.apply(self.conn, p)
        kept = self.conn.execute("SELECT value_thb, entered_by FROM snapshots WHERE account_id=? AND as_of='2026-03-26'",
                                 (stocks,)).fetchone()
        self.assertEqual((kept["value_thb"], kept["entered_by"]), (1, "best"))
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM flows").fetchone()[0], 4)   # not doubled


if __name__ == "__main__":
    unittest.main()
