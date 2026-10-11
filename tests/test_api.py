import json
import os
import sys
import tempfile
import threading
import unittest
import urllib.error
import urllib.request

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "server"))

import app  # noqa: E402
import auth  # noqa: E402
import config  # noqa: E402
import db  # noqa: E402

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None


class ApiTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.token = auth.new_token()            # generated for this run only
        cfg = dict(config.load(), data_dir=cls.tmp.name, backup_dir=os.path.join(cls.tmp.name, "backups"),
                   shop_db_path=os.path.join(cls.tmp.name, "no-shop.db"),
                   family_xlsx_path=os.path.join(cls.tmp.name, "no-family.xlsx"),
                   journal_workbook_path=os.path.join(cls.tmp.name, "journal.xlsx"),
                   agent_tokens=[{"name": "maki", "sha256": auth.hash_token(cls.token)}])
        cfg["db_path"] = os.path.join(cls.tmp.name, "t.db")
        cls.srv = app.make_server(cfg, port=0)
        cls.base = f"http://127.0.0.1:{cls.srv.server_address[1]}"
        threading.Thread(target=cls.srv.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.srv.shutdown()
        cls.srv.server_close()
        cls.tmp.cleanup()

    def call(self, method, path, body=None, token=None, cookie=None, headers=None):
        req = urllib.request.Request(self.base + path, method=method,
                                     data=None if body is None else json.dumps(body).encode("utf-8"))
        if body is not None:
            req.add_header("Content-Type", "application/json")
        if token:
            req.add_header("Authorization", f"Bearer {token}")
        if cookie:
            req.add_header("Cookie", f"{auth.UI_COOKIE}={cookie}")
        for k, v in (headers or {}).items():
            req.add_header(k, v)
        try:
            with urllib.request.urlopen(req, timeout=10) as r:
                return r.status, json.loads(r.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read().decode("utf-8"))

    def account_id(self, name):
        return next(a["id"] for a in self.call("GET", "/api/accounts")[1] if a["name"] == name)

    def test_read_needs_no_token_but_write_does(self):
        status, home = self.call("GET", "/api/home")
        self.assertEqual(status, 200)
        self.assertIn("kpis", home)
        status, out = self.call("POST", "/api/advice", {"topic": "t", "body": "b"})
        self.assertEqual(status, 401)
        status, _ = self.call("POST", "/api/advice", {"topic": "t", "body": "b"}, token="wrong-token")
        self.assertEqual(status, 401)

    def test_maki_writes_are_logged_and_can_be_undone(self):
        status, made = self.call("POST", "/api/advice", {"topic": "Risk", "body": "Risk above plan on 3 trades",
                                                          "refs": {"month": "2026-09"}}, token=self.token)
        self.assertEqual(status, 200)
        entry = next(e for e in self.call("GET", "/api/audit")[1] if e["id"] == made["audit_id"])
        self.assertEqual((entry["actor"], entry["action"], entry["tbl"]), ("maki", "insert", "advice_notes"))
        self.assertEqual(json.loads(entry["after"])["topic"], "Risk")

        status, _ = self.call("POST", f"/api/audit/{made['audit_id']}/undo", token=self.token)
        self.assertEqual(status, 200)
        self.assertFalse([n for n in self.call("GET", "/api/advice")[1] if n["id"] == made["id"]])
        status, _ = self.call("POST", f"/api/audit/{made['audit_id']}/undo", token=self.token)
        self.assertEqual(status, 400)          # already undone

    def test_maki_cannot_change_counting_rules_or_goals(self):
        acc = self.account_id("Family Statement")
        status, _ = self.call("PUT", f"/api/accounts/{acc}", {"counts_in_total": 1}, token=self.token)
        self.assertEqual(status, 403)
        status, _ = self.call("POST", "/api/goals", {"year": 2026, "name": "Stock", "bucket": "investment",
                                                      "goal_thb": 1}, token=self.token)
        self.assertEqual(status, 403)

    def test_best_writes_through_the_page_session(self):
        status, out = self.call("POST", "/api/goals", {"year": 2030, "name": "Stock", "bucket": "investment",
                                                        "goal_thb": 2000000}, cookie=auth.UI_KEY)
        self.assertEqual(status, 200)
        entry = next(e for e in self.call("GET", "/api/audit")[1] if e["id"] == out["audit_id"])
        self.assertEqual(entry["actor"], "best")
        status, _ = self.call("POST", "/api/goals", {"year": 2031, "name": "Stock", "bucket": "investment",
                                                      "goal_thb": 1}, cookie="not-the-session")
        self.assertEqual(status, 401)

    @unittest.skipIf(Workbook is None, "openpyxl not installed")
    def test_sync_makes_the_trades_match_the_workbook(self):
        from datetime import datetime
        from sources.journal import DEFAULT_SHEET

        def write(rows):
            wb = Workbook()
            ws = wb.active
            ws.title = DEFAULT_SHEET
            for n, (day, pair, win) in enumerate(rows, start=3):
                ws.cell(n, 2, datetime(2031, 5, day))
                ws.cell(n, 5, pair)
                ws.cell(n, 22, 1 if win else None)
                ws.cell(n, 25, 10.0 if win else None)
            wb.save(os.path.join(self.tmp.name, "journal.xlsx"))
        self.assertEqual(self.call("POST", "/api/journal/sync", cookie=auth.UI_KEY)[0], 404)       # no workbook yet
        self.call("POST", "/api/trades", {"traded_on": "2031-01-01", "pair": "KEEP/ME", "pnl": 1}, cookie=auth.UI_KEY)   # typed in the page
        write([(1, "BTC/USDT", True), (2, "ETH/USDT", True)])
        self.assertEqual(self.call("POST", "/api/journal/sync")[0], 401)
        status, out = self.call("POST", "/api/journal/sync", cookie=auth.UI_KEY)
        self.assertEqual((status, out["trades"], out["added"]), (200, 2, 2))
        self.assertEqual(sorted(t["pair"] for t in self.call("GET", "/api/trades")[1]), ["BTC/USDT", "ETH/USDT"])   # the typed one is gone
        write([(1, "BTC/USDT", True)])                                                                             # one row removed in Excel
        out = self.call("POST", "/api/journal/sync", cookie=auth.UI_KEY)[1]
        self.assertEqual((out["trades"], out["added"], out["removed"]), (1, 0, 1))
        self.assertEqual(len(self.call("GET", "/api/trades")[1]), 1)
        self.assertFalse(self.call("GET", "/api/trading")[1]["journal_source"]["stale"])
        self.assertEqual(self.call("POST", "/api/journal/open", token=self.token)[0], 403)                          # Maki may not open files
        for t in self.call("GET", "/api/trades")[1]:                                                                # the database is shared by the tests
            self.call("DELETE", f"/api/trades/{t['id']}", cookie=auth.UI_KEY)

    def test_a_goal_carries_its_action_plan_and_the_plan_page_loads(self):
        goal = {"year": 2044, "name": "Stock", "bucket": "investment", "goal_thb": 2000, "monthly_thb": 100, "growth_pct": 0.15, "note": "plan"}
        status, made = self.call("POST", "/api/goals", goal, cookie=auth.UI_KEY)
        self.assertEqual(status, 200)
        self.assertEqual(self.call("POST", "/api/goals", dict(goal, name="X", growth_pct=15), cookie=auth.UI_KEY)[0], 400)   # 15 is not a share
        kept = next(g for g in self.call("GET", "/api/goals?year=2044")[1] if g["id"] == made["id"])
        self.assertEqual((kept["monthly_thb"], kept["growth_pct"], kept["note"]), (100, 0.15, "plan"))
        status, plan = self.call("GET", "/api/plan")
        self.assertEqual(status, 200)
        self.assertIn("advice", plan)
        self.call("DELETE", f"/api/goals/{made['id']}", cookie=auth.UI_KEY)

    def test_a_reinvest_line_is_also_the_deposit_into_where_it_went(self):
        stocks, okx = self.account_id("Stocks"), self.account_id("OKX Trading")
        flows = lambda acc: [f for f in self.call("GET", f"/api/flows?account_id={acc}")[1] if f["note"] == "Cash flow log · Reinvest"]   # noqa: E731
        line = {"month": "2033-04", "section": "allocation", "line": "Stock", "amount_thb": 20000}
        status, made = self.call("POST", "/api/cashflow", line, cookie=auth.UI_KEY)
        self.assertEqual((status, made["deposit_to"]), (200, "Stocks"))
        self.assertEqual([(f["on_date"], f["amount_thb"], f["kind"]) for f in flows(stocks)], [("2033-04-30", 20000, "deposit")])
        self.call("PUT", f"/api/cashflow/{made['id']}", {"amount_thb": 25000}, cookie=auth.UI_KEY)                      # edited: the deposit follows
        self.assertEqual([f["amount_thb"] for f in flows(stocks)], [25000])
        self.call("PUT", f"/api/cashflow/{made['id']}", {"line": "Trading"}, cookie=auth.UI_KEY)                         # moved to another place
        self.assertEqual((len(flows(stocks)), [f["amount_thb"] for f in flows(okx)]), (0, [25000]))
        self.call("PUT", f"/api/cashflow/{made['id']}", {"line": "Saving Cash"}, cookie=auth.UI_KEY)                     # the ledger: nothing to mirror
        self.assertEqual((len(flows(okx)), len(flows(stocks))), (0, 0))
        self.call("PUT", f"/api/cashflow/{made['id']}", {"line": "Stock"}, cookie=auth.UI_KEY)
        self.assertEqual(len(flows(stocks)), 1)
        self.call("DELETE", f"/api/cashflow/{made['id']}", cookie=auth.UI_KEY)                                            # deleted: so is the deposit
        self.assertEqual(len(flows(stocks)), 0)
        income = self.call("POST", "/api/cashflow", {"month": "2033-04", "section": "income", "line": "Salary", "amount_thb": 5}, cookie=auth.UI_KEY)[1]
        self.assertIsNone(income["deposit_to"])
        self.call("DELETE", f"/api/cashflow/{income['id']}", cookie=auth.UI_KEY)

    def test_goal_bars_give_a_years_goals_against_where_they_stand(self):
        status, out = self.call("GET", "/api/goal-bars?year=2099")                 # a year with no Goal: nothing to draw
        self.assertEqual(status, 200)
        self.assertEqual(out["bars"], {})
        status, out = self.call("GET", "/api/goal-bars")
        self.assertEqual(status, 200)
        self.assertIn("bars", out)

    def test_plan_flow_rows_are_best_only_and_leave_the_log_alone(self):
        rows = [{"group": "Active", "name": "Salary", "amount": 120000}, {"group": "Investment", "name": "Stock", "amount": 5000}]
        body = {"box": "income", "rows": rows}
        self.assertEqual(self.call("PUT", "/api/plan/flow", body, token=self.token)[0], 403)
        self.assertEqual(self.call("PUT", "/api/plan/flow", {"box": "nope", "rows": rows}, cookie=auth.UI_KEY)[0], 400)
        self.assertEqual(self.call("PUT", "/api/plan/flow", {"box": "income", "rows": [{"group": "Fixed costs", "name": "x", "amount": 1}]}, cookie=auth.UI_KEY)[0], 400)
        self.assertEqual(self.call("PUT", "/api/plan/flow", {"box": "income", "rows": [{"group": "Active", "name": "x", "amount": -1}]}, cookie=auth.UI_KEY)[0], 400)
        status, saved = self.call("PUT", "/api/plan/flow", body, cookie=auth.UI_KEY)
        self.assertEqual((status, saved["income"][0]["amount"]), (200, 120000.0))
        self.assertEqual(self.call("PUT", "/api/plan/flow", {"box": "income", "rows": None}, cookie=auth.UI_KEY), (200, {}))

    def test_plan_debts_and_reinvest_are_best_only(self):
        debt = {"name": "House debt", "balance": 2156326, "as_of": "2026-10", "rate": 0.031, "payment": 35000, "line": "House debt"}
        self.assertEqual(self.call("PUT", "/api/plan/debts", {"debts": [debt]}, token=self.token)[0], 403)
        self.assertEqual(self.call("PUT", "/api/plan/debts", {"debts": [dict(debt, rate=3.1)]}, cookie=auth.UI_KEY)[0], 400)      # 3.1 is not a share
        self.assertEqual(self.call("PUT", "/api/plan/debts", {"debts": [dict(debt, payment=0)]}, cookie=auth.UI_KEY)[0], 400)
        self.assertEqual(self.call("PUT", "/api/plan/debts", {"debts": [dict(debt, as_of="Oct 2026")]}, cookie=auth.UI_KEY)[0], 400)
        self.assertEqual(self.call("PUT", "/api/plan/debts", {"debts": [debt, debt]}, cookie=auth.UI_KEY)[0], 400)
        self.assertEqual(self.call("PUT", "/api/plan/debts", {"debts": [debt]}, cookie=auth.UI_KEY), (200, [debt]))
        self.assertEqual(self.call("PUT", "/api/plan/debts", {"debts": []}, cookie=auth.UI_KEY), (200, []))
        self.assertEqual(self.call("PUT", "/api/plan/reinvest", {"stock": 1}, token=self.token)[0], 403)
        self.assertEqual(self.call("PUT", "/api/plan/reinvest", {"stock": -1}, cookie=auth.UI_KEY)[0], 400)
        self.assertEqual(self.call("PUT", "/api/plan/reinvest", {"stock": 15000, "trading": None, "freed": True}, cookie=auth.UI_KEY), (200, {"freed": True, "stock": 15000.0}))
        self.assertEqual(self.call("PUT", "/api/plan/reinvest", {}, cookie=auth.UI_KEY), (200, {"freed": False}))

    def test_plan_yields_are_best_only_shares(self):
        self.assertEqual(self.call("PUT", "/api/plan/yields", {"stock": 0.1, "trading": 0.3}, token=self.token)[0], 403)
        self.assertEqual(self.call("PUT", "/api/plan/yields", {"stock": 15, "trading": 0.5}, cookie=auth.UI_KEY)[0], 400)      # 15 is not a share
        self.assertEqual(self.call("PUT", "/api/plan/yields", {"stock": 0.12, "trading": 0.4}, cookie=auth.UI_KEY), (200, {"stock": 0.12, "trading": 0.4}))
        self.call("PUT", "/api/plan/yields", {"stock": 0.15, "trading": 0.5}, cookie=auth.UI_KEY)

    def test_cash_flow_takes_an_expense_section(self):
        status, out = self.call("POST", "/api/cashflow", {"month": "2031-02", "section": "expense", "line": "Test", "amount_thb": 100}, cookie=auth.UI_KEY)
        self.assertEqual(status, 200)
        self.assertEqual(self.call("POST", "/api/cashflow", {"month": "2031-02", "section": "bogus", "line": "x", "amount_thb": 1}, cookie=auth.UI_KEY)[0], 400)
        self.call("DELETE", f"/api/cashflow/{out['id']}", cookie=auth.UI_KEY)

    def test_a_wealth_card_is_deleted_with_its_rows_and_only_by_best(self):
        status, card = self.call("POST", "/api/accounts", {"name": "Temp card", "bucket": "wealth", "currency": "THB",
                                                           "opening_value": 100}, cookie=auth.UI_KEY)
        self.assertEqual(status, 200)
        self.assertEqual(self.call("DELETE", f"/api/accounts/{card['id']}")[0], 401)
        status, out = self.call("DELETE", f"/api/accounts/{card['id']}", cookie=auth.UI_KEY)
        self.assertEqual((status, out["removed"]["snapshots"]), (200, 1))
        self.assertNotIn("Temp card", [a["name"] for a in self.call("GET", "/api/accounts")[1]])
        self.assertEqual(self.call("DELETE", f"/api/accounts/{card['id']}", cookie=auth.UI_KEY)[0], 404)

    def test_snapshot_needs_fx_rate_for_foreign_currency_and_rejects_duplicates(self):
        acc = self.account_id("OKX Trading")
        body = {"account_id": acc, "as_of": "2026-10-31", "value_native": 2000}
        status, out = self.call("POST", "/api/snapshots", body, token=self.token)
        self.assertEqual(status, 400)
        self.assertIn("fx_rate", out["error"])
        status, made = self.call("POST", "/api/snapshots", dict(body, fx_rate=33), token=self.token)
        self.assertEqual(status, 200)
        status, _ = self.call("POST", "/api/snapshots", dict(body, fx_rate=33), token=self.token)
        self.assertEqual(status, 409)
        snap = next(s for s in self.call("GET", f"/api/snapshots?account_id={acc}")[1] if s["id"] == made["id"])
        self.assertEqual(snap["value_thb"], 66000)
        self.assertEqual(snap["entered_by"], "maki")

    def test_family_statement_card_takes_no_snapshots(self):
        acc = self.account_id("Family Statement")
        status, _ = self.call("POST", "/api/snapshots", {"account_id": acc, "as_of": "2026-10-31",
                                                          "value_native": 1}, token=self.token)
        self.assertEqual(status, 400)

    def test_new_wealth_card_opening_value_becomes_a_flow(self):
        status, out = self.call("POST", "/api/accounts", {"name": "Gold", "bucket": "wealth", "opening_value": 40000,
                                                           "opening_date": "2026-10-05"}, cookie=auth.UI_KEY)
        self.assertEqual(status, 200)
        flows = self.call("GET", f"/api/flows?account_id={out['id']}")[1]
        self.assertEqual([(f["kind"], f["amount_thb"]) for f in flows], [("opening", 40000)])

    def test_every_page_has_a_resource_maki_can_add_to_and_edit(self):
        bodies = {
            "trades": {"traded_on": "2026-10-05", "pair": "BTC/USDT", "side": "L", "outcome": "win", "pnl": 120.5},
            "benchmarks": {"name": "vt", "as_of": "2026-09-30", "value": 131.2},
            "cashflow": {"month": "2026-10", "section": "income", "line": "Salary", "amount_thb": 80000},
            "assets": {"name": "House", "asset_class": "Property", "value_thb": 3500000, "valued_on": "2026-10-01",
                       "debt_thb": 2100000},
            "upcoming": {"due_date": "2026-11-01", "item": "Insurance", "amount_thb": 18500},
            "projects": {"name": "Trading robot", "progress": 0.4},
        }
        for resource, body in bodies.items():
            status, made = self.call("POST", f"/api/{resource}", body, token=self.token)
            self.assertEqual(status, 200, (resource, made))
            status, _ = self.call("PUT", f"/api/{resource}/{made['id']}", {"note": "edited"} if resource in
                                  ("trades", "cashflow", "assets") else {}, token=self.token)
            self.assertEqual(status, 200, resource)
        self.assertEqual(self.call("GET", "/api/benchmarks?name=VT")[1][0]["name"], "VT")
        worth = self.call("GET", "/api/networth")[1]
        self.assertEqual(worth["net_worth_thb"], worth["portfolio_thb"] + 3500000 - 2100000)
        status, out = self.call("POST", "/api/cashflow", {"month": "2026-13", "section": "income", "line": "x",
                                                           "amount_thb": 1}, token=self.token)
        self.assertEqual(status, 400)

    def test_project_delete_is_logged_and_can_be_undone(self):
        _, made = self.call("POST", "/api/projects", {"name": "House design", "progress": 0.5}, cookie=auth.UI_KEY)
        status, gone = self.call("DELETE", f"/api/projects/{made['id']}", cookie=auth.UI_KEY)
        self.assertEqual(status, 200)
        self.assertFalse([p for p in self.call("GET", "/api/projects")[1] if p["id"] == made["id"]])
        self.call("POST", f"/api/audit/{gone['audit_id']}/undo", cookie=auth.UI_KEY)
        self.assertTrue([p for p in self.call("GET", "/api/projects")[1] if p["id"] == made["id"]])

    def test_home_reports_the_last_finished_month(self):
        from datetime import date
        acc = self.account_id("Stocks")
        this_month = date.today().strftime("%Y-%m")
        for as_of, value in (("2020-01-31", 100000), ("2020-02-29", 110000), (this_month + "-01", 999999)):
            status, _ = self.call("POST", "/api/snapshots", {"account_id": acc, "as_of": as_of, "value_native": value},
                                  cookie=auth.UI_KEY)
            self.assertEqual(status, 200)
        home = self.call("GET", "/api/home")[1]
        self.assertLess(home["kpis"]["month"], this_month)
        self.assertNotEqual(home["kpis"]["investment"], 999999)
        self.assertEqual(self.call("GET", "/api/monthly")[1][-1]["month"], this_month)   # still listed in full

    def test_only_best_changes_the_profile_photo(self):
        photo = "data:image/jpeg;base64,/9j/4AAQSkZJRg=="
        self.assertEqual(self.call("GET", "/api/avatar")[1], {"image": None})
        self.assertEqual(self.call("PUT", "/api/avatar", {"image": photo}, token=self.token)[0], 403)   # Maki may not
        self.assertEqual(self.call("PUT", "/api/avatar", {"image": photo})[0], 401)                      # no session at all
        self.assertEqual(self.call("PUT", "/api/avatar", {"image": "data:text/html;base64,AAAA"}, cookie=auth.UI_KEY)[0], 400)
        self.assertEqual(self.call("PUT", "/api/avatar", {"image": photo}, cookie=auth.UI_KEY)[0], 200)
        self.assertEqual(self.call("GET", "/api/avatar")[1], {"image": photo})
        self.assertEqual(self.call("PUT", "/api/avatar", {"image": None}, cookie=auth.UI_KEY)[0], 200)
        self.assertEqual(self.call("GET", "/api/avatar")[1], {"image": None})

    def test_holdings_make_the_snapshot_value_and_a_plain_total_clears_them(self):
        stocks = self.account_id("Stocks")
        hold = [{"market": "us", "stock": "VOO", "qty": 2, "value_thb": 30000}, {"market": "SET", "stock": "PTT", "qty": 100, "value_thb": 12500.5}]
        status, made = self.call("POST", "/api/snapshots", {"account_id": stocks, "as_of": "2031-03-31", "holdings": hold}, cookie=auth.UI_KEY)
        self.assertEqual(status, 200, made)
        row = next(r for r in self.call("GET", "/api/snapshots?account_id=%d" % stocks)[1] if r["as_of"] == "2031-03-31")
        self.assertEqual(row["value_thb"], 42500.5)
        self.assertEqual(json.loads(row["holdings"])[0]["market"], "US")
        self.assertEqual(self.call("POST", "/api/snapshots", {"account_id": stocks, "as_of": "2031-04-30", "holdings": [{"market": "US", "value_thb": 0}]}, cookie=auth.UI_KEY)[0], 400)
        self.assertEqual(self.call("POST", "/api/snapshots", {"account_id": stocks, "as_of": "2031-04-30", "holdings": []}, cookie=auth.UI_KEY)[0], 400)
        usdt = self.account_id("OKX Trading")
        self.assertEqual(self.call("POST", "/api/snapshots", {"account_id": usdt, "as_of": "2031-04-30", "holdings": hold}, cookie=auth.UI_KEY)[0], 400)
        status, _ = self.call("PUT", "/api/snapshots/%d" % made["id"], {"value_native": 50000}, cookie=auth.UI_KEY)
        self.assertEqual(status, 200)
        row = next(r for r in self.call("GET", "/api/snapshots?account_id=%d" % stocks)[1] if r["as_of"] == "2031-03-31")
        self.assertEqual((row["value_thb"], row["holdings"]), (50000, None))
        self.call("DELETE", "/api/snapshots/%d" % made["id"], cookie=auth.UI_KEY)       # the database is shared by the tests

    def test_investment_page_carries_stats_growth_and_markets(self):
        inv = self.call("GET", "/api/investment")[1]
        for key in ("stats", "growth_year", "growth", "markets", "benchmarks"):
            self.assertIn(key, inv)
        trading = self.call("GET", "/api/trading")[1]
        self.assertIn("stats", trading)
        self.assertIn("benchmarks", trading)
        if trading["stats"]:
            self.assertIn("fx", trading["stats"])

    def test_trade_takes_the_sheet_columns_and_refuses_a_link_that_is_not_one(self):
        status, made = self.call("POST", "/api/trades", {"traded_on": "2026-06-01", "pair": "BTC/USDT", "poi": "EMA 26", "signal": "Engulf",
                                                          "divergence": "Yes", "fib": "Fibo 0.68", "exit_a": "100 of A", "rr": 3,
                                                          "pic_entry": "https://www.tradingview.com/x/abc/", "outcome": "win", "pnl": 120.5}, token=self.token)
        self.assertEqual(status, 200, made)
        sheet = self.call("GET", "/api/trading")[1]["sheet"]
        row = next(t for t in sheet["trades"] if t["id"] == made["id"])
        self.assertEqual((row["poi"], row["exit_a"], row["rr"], row["profit"]), ("EMA 26", "100 of A", 3, 120.5))
        self.assertEqual(self.call("PUT", "/api/trades/%d" % made["id"], {"pic_exit": "javascript:alert(1)"}, token=self.token)[0], 400)
        self.call("DELETE", "/api/trades/%d" % made["id"], token=self.token)

    def test_shop_months_are_typed_in_the_dashboard_and_best_only(self):
        lines = [{"section": "revenue", "line": "Lazada", "amount": 1000}, {"section": "revenue", "line": "Shopee", "amount": "2,5"},
                 {"section": "sga", "line": "Laz Fees", "amount": 100}, {"section": "cogs", "line": "Alipay", "amount": ""}]
        save = lambda body, **kw: self.call("PUT", "/api/shop/month", body, **kw)          # noqa: E731
        self.assertEqual(save({"year": 2031, "month": 5, "entries": lines}, token=self.token)[0], 403)       # Maki may not
        self.assertEqual(save({"year": 2031, "month": 5, "entries": lines})[0], 401)
        bad = [{"section": "revenue", "line": "Lazada", "amount": 1000}, {"section": "revenue", "line": "Lazada", "amount": 5}]
        self.assertEqual(save({"year": 2031, "month": 5, "entries": bad}, cookie=auth.UI_KEY)[0], 400)       # same line twice
        self.assertEqual(save({"year": 2031, "month": 13, "entries": bad[:1]}, cookie=auth.UI_KEY)[0], 400)
        ok = [lines[0], {"section": "revenue", "line": "Shopee", "amount": 2500.5}, lines[2], lines[3]]
        status, made = save({"year": 2031, "month": 5, "entries": ok}, cookie=auth.UI_KEY)
        self.assertEqual(status, 200, made)
        year = next(y for y in self.call("GET", "/api/shop/finance")[1]["years"] if y["year"] == 2031)
        m = year["months"][0]
        self.assertEqual((m["month"], m["revenue"], m["sga"], m["after_fees"], m["has_cogs"]), (5, 3500.5, 100, 3400.5, False))
        self.assertEqual(m["net_profit"], 3400.5)                              # as the sheet does it: no cost of goods counts as 0
        self.assertIsNone(m["dividend_pool"])                                  # no Payout % typed for the month
        self.assertEqual(save({"year": 2031, "month": 5, "entries": [lines[0]]}, cookie=auth.UI_KEY)[0], 200)      # replace the month
        year = next(y for y in self.call("GET", "/api/shop/finance")[1]["years"] if y["year"] == 2031)
        self.assertEqual(year["months"][0]["revenue"], 1000)
        self.assertTrue(any(r["month"] == "2031-05" and r["revenue"] == 1000 for r in self.call("GET", "/api/shop")[1]["months"]))
        # no entries removes the month, and the removal is in the history so it can be undone
        status, gone = save({"year": 2031, "month": 5, "entries": []}, cookie=auth.UI_KEY)
        self.assertEqual((status, gone["deleted"]), (200, True))
        self.assertFalse(any(y["year"] == 2031 for y in self.call("GET", "/api/shop/finance")[1]["years"]))
        self.assertEqual(self.call("POST", "/api/audit/%d/undo" % gone["audit_id"], cookie=auth.UI_KEY)[0], 200)
        self.assertTrue(any(y["year"] == 2031 for y in self.call("GET", "/api/shop/finance")[1]["years"]))
        save({"year": 2031, "month": 5, "entries": []}, cookie=auth.UI_KEY)        # the database is shared by the tests

    def test_dividend_follows_the_sheet_rule_and_payout_is_a_percentage(self):
        save = lambda entries: self.call("PUT", "/api/shop/month", {"year": 2032, "month": 1, "entries": entries}, cookie=auth.UI_KEY)   # noqa: E731
        base = [{"section": "revenue", "line": "Lazada", "amount": 10000}, {"section": "sga", "line": "Laz Fees", "amount": 2000},
                {"section": "cogs", "line": "Alipay", "amount": 3000}]
        self.assertEqual(save(base + [{"section": "dividend", "line": "Payout %", "amount": 150}])[0], 400)
        self.assertEqual(save(base + [{"section": "dividend", "line": "Whatever", "amount": 10}])[0], 400)
        self.assertEqual(save(base + [{"section": "dividend", "line": "Payout %", "amount": 38}])[0], 200)
        finance = self.call("GET", "/api/shop/finance")[1]
        m = next(y for y in finance["years"] if y["year"] == 2032)["months"][0]
        self.assertEqual(m["net_profit"], 5000)                                # 10,000 - 3,000 - 2,000
        self.assertEqual(m["dividend_pool"], 3040)                             # (10,000 - 2,000) x 38%, cost of goods not taken off
        self.assertEqual((m["dividend_onn"], m["dividend_bes"]), (1824, 1216))  # 60 / 40
        self.assertEqual(finance["split_onn"], 0.6)
        self.assertGreaterEqual(finance["dividend_total"]["onn"], 1824)
        save([])                                                               # the database is shared by the tests

    def test_the_shop_apps_typed_months_are_copied_once_and_never_overwrite_edits(self):
        import sqlite3
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            shop_db = os.path.join(tmp, "shop.db")
            c = sqlite3.connect(shop_db)
            c.execute("CREATE TABLE finance_monthly (year INT, month INT, section TEXT, line TEXT, amount REAL)")
            c.executemany("INSERT INTO finance_monthly VALUES (?,?,?,?,?)", [(2030, 1, "revenue", "Lazada", 500.0), (2030, 1, "sga", "Laz Fees", 50.0), (2030, 2, "revenue", "Shopee", 70.0)])
            c.commit()
            c.close()
            conn = db.connect(os.path.join(tmp, "d.db"))
            db.init(conn)
            self.assertEqual(app.seed_shop_months(conn, shop_db), 2)
            self.assertEqual(app.shop_months(conn)["months"][0]["revenue"], 500.0)
            conn.execute("UPDATE shop_months SET entries='[]' WHERE month=1")            # Best edits a month
            conn.commit()
            self.assertEqual(app.seed_shop_months(conn, shop_db), 0)                       # a second run changes nothing
            self.assertEqual(conn.execute("SELECT entries FROM shop_months WHERE month=1").fetchone()[0], "[]")
            conn.close()
            empty = db.connect(os.path.join(tmp, "e.db"))
            db.init(empty)
            self.assertEqual(app.seed_shop_months(empty, os.path.join(tmp, "missing.db")), 0)
            empty.close()

    def test_only_best_edits_the_grade_notes(self):
        notes = self.call("GET", "/api/trading")[1]["grades"]
        self.assertEqual(sorted(notes), ["A", "B", "C"])
        new = {"A": "Plan followed, exit as planned", "B": "Plan followed, early exit", "C": "Plan broken"}
        self.assertEqual(self.call("PUT", "/api/journal/grades", new, token=self.token)[0], 403)
        self.assertEqual(self.call("PUT", "/api/journal/grades", dict(new, A="x" * 81), cookie=auth.UI_KEY)[0], 400)
        self.assertEqual(self.call("PUT", "/api/journal/grades", new, cookie=auth.UI_KEY)[0], 200)
        self.assertEqual(self.call("GET", "/api/trading")[1]["grades"], new)
        self.call("PUT", "/api/journal/grades", notes, cookie=auth.UI_KEY)

    def test_only_best_sets_the_journal_start_and_risk(self):
        self.assertEqual(self.call("PUT", "/api/journal/settings", {"start_balance": 1000, "risk_pct": 0.05}, token=self.token)[0], 403)
        self.assertEqual(self.call("PUT", "/api/journal/settings", {"start_balance": 1000, "risk_pct": 5}, cookie=auth.UI_KEY)[0], 400)
        self.assertEqual(self.call("PUT", "/api/journal/settings", {"start_balance": 1000, "risk_pct": 0.05, "start_date": "not a date"}, cookie=auth.UI_KEY)[0], 400)
        self.assertEqual(self.call("PUT", "/api/journal/settings", {"start_balance": 1000, "risk_pct": 0.05, "start_date": "2026-01-01"}, cookie=auth.UI_KEY)[0], 200)
        s = self.call("GET", "/api/trading")[1]["sheet"]["summary"]
        self.assertEqual((s["start"], s["risk_pct"]), (1000, 0.05))
        self.call("PUT", "/api/journal/settings", {"start_balance": 980, "risk_pct": 0.1}, cookie=auth.UI_KEY)

    def test_requests_from_another_origin_are_refused(self):
        status, _ = self.call("GET", "/api/home", headers={"Origin": "http://evil.example"})
        self.assertEqual(status, 403)

    def test_there_is_no_route_that_moves_money(self):
        for path in ("/api/orders", "/api/trade", "/api/transfer", "/api/withdraw"):
            status, _ = self.call("POST", path, {}, token=self.token)
            self.assertEqual(status, 404)


if __name__ == "__main__":
    unittest.main()
