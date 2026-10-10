"""Import the manual trade journal from `1.1 Trade Journal2.xlsx`.

    python scripts/import_trade_journal.py --dry-run
    python scripts/import_trade_journal.py

Reads the visible journal sheet. A trade is a row with a date in column B and a
pair in column E. Re-running updates the same trades (matched by sheet row).
Amounts stay in the journal's own currency (USDT).
"""
import argparse
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "server"))

import config  # noqa: E402
import db  # noqa: E402

DEFAULT_SHEET = "Trading Journal Fianl Test"     # spelled as in the workbook


def cell(row, letter):
    """Value by column letter(s); the journal has 29 columns (A..AC)."""
    n = 0
    for ch in letter:
        n = n * 26 + ord(ch) - 64
    return row[n - 1] if n <= len(row) else None


def clean(v):
    return str(v).strip() if v not in (None, "") else None


def link(v):
    """A chart link, or None for anything that is not one (the sheet's header rows say 'Pic')."""
    v = clean(v)
    return v if v and v.lower().startswith(("https://", "http://")) else None


def number(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def parse(path, sheet=DEFAULT_SHEET):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    trades = []
    for n, r in enumerate(wb[sheet].iter_rows(min_row=1, max_col=29, values_only=True), start=1):
        when, pair = cell(r, "B"), clean(cell(r, "E"))
        if not isinstance(when, datetime) or not pair:
            continue
        outcome = "win" if cell(r, "V") else "loss" if cell(r, "W") else "be" if cell(r, "X") else None
        # Profit (Y) and Loss (Z) are kept whatever the outcome: break-even trades carry small amounts,
        # and the sheet's own PNL is the sum of Y minus the sum of Z.
        profit, lost = number(cell(r, "Y")), number(cell(r, "Z"))
        pnl = None if profit is None and lost is None and outcome is None else (profit or 0) - abs(lost or 0)
        trades.append({"source_row": n, "traded_on": when.strftime("%Y-%m-%d"), "session": clean(cell(r, "D")),
                       "pair": pair, "side": clean(cell(r, "F")), "setup": clean(cell(r, "G")),
                       "exit_plan": clean(cell(r, "P")), "exit_actual": clean(cell(r, "Q")),
                       "grade": clean(cell(r, "U")), "outcome": outcome, "pnl": pnl,
                       "balance": number(cell(r, "AB")), "risk": number(cell(r, "AC")), "note": clean(cell(r, "T")),
                       # War Zone (I..K), Entry (L..O), Exit's "100 of A?" (R) and the three chart links (H, O, S)
                       "pic_setup": link(cell(r, "H")), "poi": clean(cell(r, "I")), "signal": clean(cell(r, "J")),
                       "divergence": clean(cell(r, "K")), "confirm": clean(cell(r, "L")), "fib_time": clean(cell(r, "M")),
                       "fib": clean(cell(r, "N")), "pic_entry": link(cell(r, "O")), "exit_a": clean(cell(r, "R")),
                       "pic_exit": link(cell(r, "S")), "rr": number(cell(r, "AA"))})
    wb.close()
    return trades


def apply(conn, trades):
    cols = list(trades[0]) if trades else []
    for t in trades:
        conn.execute(
            f"INSERT INTO trades({', '.join(cols)}) VALUES ({', '.join('?' for _ in cols)}) "
            f"ON CONFLICT(source_row) DO UPDATE SET {', '.join(c + '=excluded.' + c for c in cols if c != 'source_row')}",
            [t[c] for c in cols])
    conn.commit()


def main():
    cfg = config.load()
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--workbook", default=cfg["journal_workbook_path"])
    ap.add_argument("--sheet", default=DEFAULT_SHEET)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    trades = parse(args.workbook, args.sheet)
    graded = [t for t in trades if t["outcome"]]
    print(f"{len(trades)} trades ({len(graded)} with an outcome), "
          f"{trades[0]['traded_on'] if trades else '-'} .. {trades[-1]['traded_on'] if trades else '-'}")
    print("  win {win}  loss {loss}  be {be}  pnl {pnl:,.2f}".format(
        win=sum(t["outcome"] == "win" for t in trades), loss=sum(t["outcome"] == "loss" for t in trades),
        be=sum(t["outcome"] == "be" for t in trades), pnl=sum(t["pnl"] or 0 for t in trades)))
    if args.dry_run:
        print("dry run: nothing written")
        return
    conn = db.connect(cfg["db_path"])
    db.init(conn)
    apply(conn, trades)
    conn.close()
    print("imported.")


if __name__ == "__main__":
    main()
