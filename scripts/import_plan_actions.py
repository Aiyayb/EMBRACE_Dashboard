"""Bring the action plan of the year from the sheet "Fin Plan(Y)" of `4. Fund Management.xlsx` onto that year's Goals.

    python scripts/import_plan_actions.py --dry-run     # show what would be set
    python scripts/import_plan_actions.py               # write it

The Goals themselves (the TARGET table) come in with import_fund_management.py. This adds what the sheet's "Action Plan" block
says about each Goal of its year: the monthly reinvestment, the growth the plan counts on, and Best's own conclusion as the note.
For OnTheWall the monthly figure is the income aimed at each month. It also records Best's birth year (the sheet's year minus
the age beside it) so the Plan page can show his age.

Only empty fields are filled: anything Best has typed on the Plan page is left as it is. The block's labels are checked first,
and nothing is written if the sheet no longer looks the way this script expects. A backup of the database is made before writing.
"""
import argparse
import os
import re
import shutil
import sys
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "server"))

import config  # noqa: E402
import db  # noqa: E402

SHEET = "Fin Plan(Y)"
# Goal name -> (label cell, expected label, monthly cell, growth cell or None, note cell)
BLOCK = {
    "Stock": ("J48", "Monthly", "K48", "K49", "J50"),
    "Trading": ("L48", "Monthly", "M48", "M49", "L50"),
    "Saving Bitcoin": ("N48", "Monthly", "O48", "O49", "N50"),         # the sheet's Saving column
    "Saving Cash": ("P48", "Monthly", "Q48", None, "P50"),             # the sheet's EMC column
    "OnTheWall": ("I42", "Monthly", "I43", None, "I44"),
}


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def words(v):
    return " · ".join(part.strip() for part in str(v).splitlines() if part.strip()) if v not in (None, "") else None


def parse(ws):
    title = str(ws["F41"].value or "")
    found = re.search(r"(20\d\d)", title)
    if "Action Plan" not in title or not found:
        raise SystemExit(f"stopping: F41 should hold 'Action Plan <year>', it holds {title!r} (nothing written)")
    plan = {}
    for name, (label_cell, label, monthly, growth, note) in BLOCK.items():
        if str(ws[label_cell].value or "").strip() != label:
            raise SystemExit(f"stopping: {label_cell} should say {label!r} for {name}, it says {ws[label_cell].value!r} (nothing written)")
        plan[name] = {"monthly_thb": num(ws[monthly].value), "growth_pct": num(ws[growth].value) if growth else None, "note": words(ws[note].value)}
    age, year = num(ws["B16"].value), num(ws["C16"].value)
    return int(found.group(1)), plan, int(year - age) if age and year else None


def main():
    cfg = config.load()
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--workbook", default=cfg["fund_workbook_path"])
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    from openpyxl import load_workbook
    wb = load_workbook(args.workbook, data_only=True)
    year, plan, birth = parse(wb[SHEET])
    wb.close()

    conn = db.connect(cfg["db_path"])
    db.init(conn)
    writes = []
    for name, fields in plan.items():
        goal = conn.execute("SELECT * FROM goals WHERE year=? AND name=?", (year, name)).fetchone()
        if goal is None:
            print(f"{year} {name}: no such Goal, skipped")
            continue
        fill = {k: v for k, v in fields.items() if v is not None and goal[k] is None}
        kept = [k for k, v in fields.items() if v is not None and goal[k] is not None and goal[k] != v]
        print(f"{year} {name}: " + (", ".join(f"{k} = {v}" for k, v in fill.items()) or "nothing to set")
              + (f"   (kept as typed: {', '.join(kept)})" if kept else ""))
        if fill:
            writes.append((goal["id"], fill))
    set_birth = birth and not db.get_setting(conn, "plan_birth_year")
    if set_birth:
        print(f"birth year: {birth}")
    if args.dry_run or not (writes or set_birth):
        print("dry run: nothing written" if args.dry_run else "nothing to do")
        return
    os.makedirs(cfg["backup_dir"], exist_ok=True)
    conn.commit()
    keep = os.path.join(cfg["backup_dir"], f"dashboard-{datetime.now():%Y%m%d-%H%M%S}-pre-plan-import.db")
    shutil.copyfile(cfg["db_path"], keep)
    print("backup:", keep)
    for goal_id, fill in writes:
        conn.execute(f"UPDATE goals SET {', '.join(k + '=?' for k in fill)} WHERE id=?", [*fill.values(), goal_id])
    if set_birth:
        db.set_setting(conn, "plan_birth_year", birth)
    conn.commit()
    conn.close()
    print(f"wrote {len(writes)} Goals")


if __name__ == "__main__":
    main()
