"""Bring the shop's monthly P&L from `ONTHEWALL SHOP BS.xlsx` into the dashboard (shop_months).

    python scripts/import_shop_finance.py --dry-run     # show what would change and check it against the sheet
    python scripts/import_shop_finance.py               # write it

Reads the visible sheets "OTW Finance 24", "OTW Finance 25" and "OTW Finance 26". The year of each column comes from the
dates in the block's header row (the 25 and 26 sheets' blocks are still labelled "Year 2024"). It does three things:

  * a month the dashboard does not have yet (2024 and 2025 at first) is added in full;
  * for a month it already has, the sheet's Shopee Fees (row 38) is set when the sheet has a number, because that row was
    filled in later; every other line Best typed in the dashboard is left as it is, and any difference is only reported;
  * every month gets its "Payout %" (the dividend rate the sheet applies to revenue minus fees: 38 or 33), worked out from
    the sheet's own dividend row, for 2023 to 2026.

The numbered cost rows (Uneedcargo 1..4, Alipay 1..4) are added up into one line each; "ETC…" is its own cost line. An empty
cell is no data; a typed 0 is kept. Revenue, cost and fee totals and the dividend are checked against what the sheet itself
calculates, and nothing is written if any month disagrees. A backup of the database is made before writing.
"""
import argparse
import json
import os
import shutil
import sys
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "server"))

import config  # noqa: E402
import db  # noqa: E402

SHEETS = ("OTW Finance 24", "OTW Finance 25", "OTW Finance 26")
FIRST_YEAR, LAST_YEAR = 2024, 2026
SPLIT_ONN = 0.6
LINES = {"Lazada": ("revenue", "Lazada"), "Shopee": ("revenue", "Shopee"), "ขายตรง": ("revenue", "ขายตรง"),
         "ETC…": ("cogs", "ETC"), "Laz Transport": ("sga", "Laz Transport"), "Laz Marketing": ("sga", "Laz Marketing"),
         "Laz Fees": ("sga", "Laz Fees"), "Shopee Fees": ("sga", "Shopee Fees"), "ดอกเบี้ย": ("tax", "ดอกเบี้ย"), "ภาษี": ("tax", "ภาษี")}
TOTALS = {"รายได้รวม": "revenue", "รวมต้นทุน": "cogs", "รวมต้นทุนบริหาร": "sga"}
POOL_LABELS = ("กำไร", "กำไรจากการดำเนินงาน")      # the dividend pool row; the 26 sheet puts its label in column B
PAYOUT = ("dividend", "Payout %")
ORDER = {"revenue": 0, "cogs": 1, "sga": 2, "tax": 3, "dividend": 4}


def number(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def payout_pct(base, pool, onn, bes):
    """The rate behind a month's dividend row, or None for a month without one; checks the sheet's own 60/40 split."""
    if not base or base <= 0 or not pool:
        return None, None
    pct = round(pool / base * 100)
    problems = []
    if abs(base * pct / 100 - pool) > 0.5:
        problems.append(f"dividend pool {pool:,.2f} is not {pct}% of {base:,.2f}")
    if onn is not None and abs(pool * SPLIT_ONN - onn) > 0.5:
        problems.append("ออน is not 60% of the pool")
    if bes is not None and abs(pool * (1 - SPLIT_ONN) - bes) > 0.5:
        problems.append("เบส is not 40% of the pool")
    return float(pct), problems


def parse_sheet(ws):
    """{(year, month): {"entries": [...], "sheet_totals": {...}, "pool": (pool, onn, bes)}} for the 2024-2026 block of one sheet."""
    rows = list(ws.iter_rows(min_row=1, max_col=16, values_only=True))
    header = next((i for i, r in enumerate(rows)
                   if str(r[1] or "").startswith("Year") and isinstance(r[3], datetime) and FIRST_YEAR <= r[3].year <= LAST_YEAR), None)
    if header is None:
        return {}
    months = {c: (rows[header][c].year, rows[header][c].month) for c in range(3, 15) if isinstance(rows[header][c], datetime)}
    out = {ym: {"cells": {}, "sheet_totals": {}, "pool": [None, None, None]} for ym in months.values()}
    dividend_rows = []
    for r in rows[header + 1:]:
        label = str(r[2] or "").strip()
        b_label = str(r[1] or "").strip()
        if label.startswith("Year") or b_label.startswith("Year"):
            break
        is_pool = label in POOL_LABELS or b_label in POOL_LABELS
        if is_pool or b_label.startswith("ปันผล") or (dividend_rows and label in ("ออน", "เบส")):
            if is_pool and dividend_rows:
                dividend_rows = []                                 # a later "profit" row starts the dividend block again
            dividend_rows.append((label or b_label, r))
            continue
        if label in TOTALS:
            for c, ym in months.items():
                out[ym]["sheet_totals"][TOTALS[label]] = number(r[c])
            continue
        base = label.rsplit(" ", 1)[0] if label.rsplit(" ", 1)[-1].isdigit() else label
        target = LINES.get(label) or (("cogs", base) if base in ("Uneedcargo", "Alipay") else None)
        if not target:
            continue
        for c, ym in months.items():
            v = number(r[c])
            if v is not None:
                out[ym]["cells"][target] = out[ym]["cells"].get(target, 0.0) + v        # the numbered cost rows add up to one line
    if len(dividend_rows) >= 3:                                                           # pool, ออน, เบส
        for k in range(3):
            for c, ym in months.items():
                out[ym]["pool"][k] = number(dividend_rows[k][1][c])
    result = {}
    for ym, d in out.items():
        entries = [{"section": s, "line": l, "amount": round(v, 2)} for (s, l), v in d["cells"].items()]
        entries.sort(key=lambda e: ORDER[e["section"]])
        result[ym] = {"entries": entries, "sheet_totals": d["sheet_totals"], "pool": tuple(d["pool"])}
    return result


def parse_2023(ws):
    """Payout % of each 2023 month: the sheet's 2023 block has revenue, a profit row (revenue x rate) and the two shares."""
    rows = list(ws.iter_rows(min_row=1, max_col=16, values_only=True))
    header = next((i for i, r in enumerate(rows) if str(r[1] or "").startswith("Year") and isinstance(r[3], datetime) and r[3].year == 2023), None)
    if header is None:
        return {}
    cols = {c: rows[header][c].month for c in range(3, 15) if isinstance(rows[header][c], datetime)}
    found = {}
    for r in rows[header + 1:header + 12]:
        found[str(r[2] or "").strip() + ("#" + str(r[1]) if r[1] else "")] = r
    revenue, profit = next(v for k, v in found.items() if k.startswith("รายได้รวม")), next(v for k, v in found.items() if k.startswith("กำไร"))
    onn, bes = next(v for k, v in found.items() if k.startswith("ออน")), next(v for k, v in found.items() if k.startswith("เบส"))
    out = {}
    for c, month in cols.items():
        pct, problems = payout_pct(number(revenue[c]), number(profit[c]), number(onn[c]), number(bes[c]))
        if pct is not None:
            out[(2023, month)] = {"pct": pct, "problems": problems}
    return out


def check(month):
    """Differences between what was read and the totals and dividend the sheet calculates for the same month."""
    sums = {}
    for e in month["entries"]:
        sums[e["section"]] = sums.get(e["section"], 0.0) + e["amount"]
    bad = [f"{k}: read {sums.get(k, 0):,.2f}, sheet {t:,.2f}" for k, t in month["sheet_totals"].items()
           if t is not None and abs(sums.get(k, 0.0) - t) > 0.5]
    pool, onn, bes = month["pool"]
    pct, problems = payout_pct(sums.get("revenue", 0.0) - sums.get("sga", 0.0), pool, onn, bes)
    month["pct"] = pct
    return bad + (problems or [])


def main():
    cfg = config.load()
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--workbook", default=os.path.join(config.EMBRACE_OS_DIR, "EMBRACE SHOP", "ONTHEWALL", "ONTHEWALL SHOP BS.xlsx"))
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    from openpyxl import load_workbook
    wb = load_workbook(args.workbook, data_only=True, read_only=False)
    found = {}
    for name in SHEETS:
        found.update(parse_sheet(wb[name]))
    old_years = parse_2023(wb[SHEETS[0]])
    wb.close()

    problems = 0
    for ym, m in sorted(found.items()):
        diff = check(m)
        problems += bool(diff)
        if diff:
            print(f"{ym[0]}-{ym[1]:02d}  MISMATCH: " + "; ".join(diff))
    for ym, m in sorted(old_years.items()):
        problems += bool(m["problems"])
        if m["problems"]:
            print(f"{ym[0]}-{ym[1]:02d}  MISMATCH: " + "; ".join(m["problems"]))
    print(f"{len(found) + len(old_years)} months read, {problems} that disagree with the sheet's own totals or dividend")
    if problems:
        print("stopping: fix the disagreement first (nothing written)")
        sys.exit(1)

    conn = db.connect(cfg["db_path"])
    db.init(conn)
    have = {(r["year"], r["month"]): json.loads(r["entries"]) for r in conn.execute("SELECT year, month, entries FROM shop_months")}
    writes, notes = {}, []
    for ym, m in sorted(found.items()):
        if m["pct"] is None and not m["entries"]:
            continue
        sheet = {(e["section"], e["line"]): e["amount"] for e in m["entries"]}
        if ym not in have:
            if not any(e["section"] == "revenue" for e in m["entries"]):
                continue                                       # nothing typed for that month in the sheet (e.g. Oct-Dec 2026)
            entries = list(m["entries"]) + ([{"section": "dividend", "line": "Payout %", "amount": m["pct"]}] if m["pct"] else [])
            writes[ym] = ("add", entries)
            continue
        current = {(e["section"], e["line"]): e for e in have[ym]}
        changed = []
        shopee = sheet.get(("sga", "Shopee Fees"))
        if shopee is not None and current.get(("sga", "Shopee Fees"), {}).get("amount") != shopee:
            changed.append(f"Shopee Fees {current.get(('sga', 'Shopee Fees'), {}).get('amount')} -> {shopee}")
            current[("sga", "Shopee Fees")] = {"section": "sga", "line": "Shopee Fees", "amount": shopee}
        if m["pct"] is not None and current.get(PAYOUT, {}).get("amount") != m["pct"]:
            changed.append(f"Payout % -> {m['pct']:g}")
            current[PAYOUT] = {"section": "dividend", "line": "Payout %", "amount": m["pct"]}
        for key, amount in sheet.items():                       # reported, never applied
            if key != ("sga", "Shopee Fees") and key in current and abs(current[key]["amount"] - amount) > 0.01:
                notes.append(f"{ym[0]}-{ym[1]:02d} {key[1]}: dashboard {current[key]['amount']:,.2f}, sheet {amount:,.2f} (left as it is)")
        if changed:
            writes[ym] = ("update " + ", ".join(changed), sorted(current.values(), key=lambda e: ORDER[e["section"]]))
    for ym, pct in sorted(((k, v["pct"]) for k, v in old_years.items())):
        entries = have.get(ym)
        if entries is None:
            continue
        current = {(e["section"], e["line"]): e for e in entries}
        if current.get(PAYOUT, {}).get("amount") != pct:
            current[PAYOUT] = {"section": "dividend", "line": "Payout %", "amount": pct}
            writes[ym] = (f"update Payout % -> {pct:g}", sorted(current.values(), key=lambda e: ORDER[e["section"]]))
    for ym, (what, _) in sorted(writes.items()):
        print(f"{ym[0]}-{ym[1]:02d}  {what}")
    for n in notes:
        print("note:", n)
    print(f"{len(writes)} months to write")
    if args.dry_run or not writes:
        print("dry run: nothing written" if args.dry_run else "nothing to do")
        return
    os.makedirs(cfg["backup_dir"], exist_ok=True)
    conn.commit()
    keep = os.path.join(cfg["backup_dir"], f"dashboard-{datetime.now():%Y%m%d-%H%M%S}-pre-shop-import.db")
    shutil.copyfile(cfg["db_path"], keep)
    print("backup:", keep)
    for (year, month), (_, entries) in sorted(writes.items()):
        conn.execute("INSERT INTO shop_months(year, month, entries, entered_by, created_at) VALUES (?,?,?,?,?) "
                     "ON CONFLICT(year, month) DO UPDATE SET entries=excluded.entries",
                     (year, month, json.dumps(entries, ensure_ascii=False), "import", db.now()))
    conn.commit()
    conn.close()
    print(f"wrote {len(writes)} months")


if __name__ == "__main__":
    main()
